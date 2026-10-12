"""Shared helpers: settings, paths, and the Excel claims register."""
from __future__ import annotations

import csv
import datetime as dt
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
POSTS = ROOT / "posts"
PROMPTS = ROOT / "prompts"
XLSX = DATA / "claims.xlsx"

LABELS = ["PROVEN", "ALLEGED", "DISPUTED", "CLEARED"]
POST_STATUSES = ["Backlog", "Ready", "Drafted", "Published", "Update-needed"]

SHEETS = {
    "Claims": [
        "claim_id", "event_date", "title", "summary", "parties", "category",
        "status_label", "govt_link", "adani_response", "source_urls",
        "verified", "verified_by", "last_verified", "post_status",
        "published_label", "post_id", "notes",
    ],
    "Inbox": ["found_at", "feed", "title", "url", "published", "triage", "claim_id"],
    "PostLog": ["post_id", "claim_id", "kind", "created_at", "status_label",
                "published_at", "ig_media_id", "permalink"],
}

HELP = {
    "status_label": "PROVEN = court/regulator finding; ALLEGED = reported or under investigation; "
                    "DISPUTED = denied by the group; CLEARED = dismissed by an authority",
    "source_urls": "One URL per line. Primary sources first (court order, SEBI, DOJ, CAG).",
    "verified": "Set to Yes only after a human has opened every source on the day.",
    "post_status": "Ready = queue for drafting. The pipeline sets Drafted / Published / Update-needed.",
}


def settings() -> dict:
    return yaml.safe_load((ROOT / "config" / "settings.yml").read_text(encoding="utf-8"))


def today() -> str:
    return dt.date.today().isoformat()


def now() -> str:
    return dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%d %H:%M UTC")


# ---------- Excel ----------

def create_workbook(path: Path = XLSX) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="101318")
    for name, cols in SHEETS.items():
        ws = wb.create_sheet(name)
        ws.append(cols)
        for i, c in enumerate(cols, 1):
            cell = ws.cell(row=1, column=i)
            cell.font = Font(bold=True, color="F5B700")
            cell.fill = head_fill
            ws.column_dimensions[cell.column_letter].width = 42 if c in (
                "summary", "source_urls", "adani_response", "title", "govt_link", "url") else 16
        ws.freeze_panes = "A2"
    claims = wb["Claims"]
    def dv(col: str, values: list[str]):
        v = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
        idx = SHEETS["Claims"].index(col) + 1
        letter = claims.cell(row=1, column=idx).column_letter
        v.add(f"{letter}2:{letter}2000")
        claims.add_data_validation(v)
    dv("status_label", LABELS)
    dv("verified", ["Yes", "No"])
    dv("post_status", POST_STATUSES)
    dv("published_label", LABELS)
    guide = wb.create_sheet("Guide")
    guide.append(["Column", "Meaning"])
    for k, v in HELP.items():
        guide.append([k, v])
    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["B"].width = 110
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def read_sheet(name: str) -> list[dict]:
    wb = load_workbook(XLSX)
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    head = rows[0]
    out = []
    for r in rows[1:]:
        if r is None or all(v in (None, "") for v in r):
            continue
        out.append({h: ("" if v is None else v) for h, v in zip(head, r)})
    return out


def write_sheet(name: str, records: list[dict]) -> None:
    """Rewrite data rows of a sheet, keeping header, styling and validations."""
    wb = load_workbook(XLSX)
    ws = wb[name]
    cols = SHEETS[name]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    wrap = Alignment(wrap_text=True, vertical="top")
    for rec in records:
        ws.append([rec.get(c, "") for c in cols])
        for cell in ws[ws.max_row]:
            cell.alignment = wrap
    wb.save(XLSX)
    export_csv()


def export_csv() -> None:
    """CSV mirrors so pull-request diffs of the Excel file are readable."""
    for name, cols in SHEETS.items():
        rows = read_sheet(name)
        with open(DATA / f"{name.lower()}.csv", "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=cols)
            w.writeheader()
            for r in rows:
                w.writerow({c: str(r.get(c, "")) for c in cols})


def source_list(claim: dict) -> list[str]:
    return [u.strip() for u in str(claim.get("source_urls", "")).splitlines() if u.strip().startswith("http")]
